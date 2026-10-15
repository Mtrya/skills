"""Exercise the real command-line tools and inspect their PPTX outputs."""

import os
from pathlib import Path
import subprocess
import sys

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
import pytest
import yaml


SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'


def make_deck(root, size=(960, 540), elements=None, **extra):
    root.mkdir(exist_ok=True)
    (root / 'pages').mkdir(exist_ok=True)
    deck = {'version': 'v2', 'size': list(size), 'pages': ['pages/one.page'], **extra}
    (root / 'deck.pptd').write_text(yaml.safe_dump(deck), encoding='utf-8')
    (root / 'pages/one.page').write_text(yaml.safe_dump({'elements': elements or []}), encoding='utf-8')
    return root


def run_script(name, root, *args):
    env = {**os.environ, 'PYTHONDONTWRITEBYTECODE': '1', 'MPLCONFIGDIR': str(root / 'mpl-cache')}
    return subprocess.run([sys.executable, str(SCRIPTS / name), str(root), *args], env=env, capture_output=True, text=True)


@pytest.mark.parametrize('size', [(960, 540), (960, 720), (960, 1280), (1080, 1080)])
def test_canvas_and_shapes_keep_source_proportions(tmp_path, size):
    root = make_deck(tmp_path, size, [{'elementId': 'circle', 'elementType': 'shape', 'bounds': [80, 80, 160, 160], 'shapeName': 'ellipse', 'fill': {'type': 'solid', 'color': '#fff'}}])
    output = root / 'result.pptx'
    result = run_script('pptd_to_pptx.py', root, '-o', str(output))
    assert result.returncode == 0, result.stderr
    deck = Presentation(output)
    assert deck.slide_width / deck.slide_height == pytest.approx(size[0] / size[1])
    circle = deck.slides[0].shapes[0]
    assert circle.width == circle.height
    assert circle.left / deck.slide_width == pytest.approx(80 / size[0])
    assert circle.auto_shape_type == MSO_SHAPE.OVAL
    assert str(circle.fill.fore_color.rgb) == 'FFFFFF'


def test_supported_content_survives_export(tmp_path):
    elements = [
        {'elementId': 'text', 'elementType': 'text', 'bounds': [20, 20, 600, 80], 'content': {'text': 'Before <strong>bold <em>both</em></strong> after &amp; more', 'textStyle': '$body'}},
        {'elementId': 'photo', 'elementType': 'image', 'bounds': [20, 120, 200, 100], 'src': 'sample.png', 'fit': {'mode': 'fill'}},
        {'elementId': 'table', 'elementType': 'table', 'bounds': [300, 120, 400, 180], 'columnWidths': [0.5, 0.5], 'rowHeights': [0.5, 0.5], 'rows': [[{'text': 'Merged', 'rowSpan': 2}, {'text': 'Top'}], [{'text': 'Bottom'}]]},
        {'elementId': 'chart', 'elementType': 'chart', 'bounds': [20, 330, 400, 180], 'series': [{'type': 'bar', 'name': 'Revenue', 'data': [2, 5]}], 'xAxis': {'type': 'category', 'data': ['A', 'B']}, 'yAxis': {'type': 'value', 'name': 'USD'}, 'legend': {'show': True}},
    ]
    root = make_deck(tmp_path, elements=elements, theme={'colors': {'ink': '#123456'}, 'textStyles': {'body': {'fontSize': 24, 'color': '$ink'}}})
    Image.new('RGB', (200, 100), 'red').save(root / 'sample.png')
    checked = run_script('pptd_check.py', root)
    assert checked.returncode == 0, checked.stdout + checked.stderr
    output = root / 'result.pptx'
    result = run_script('pptd_to_pptx.py', root, '-o', str(output))
    assert result.returncode == 0, result.stderr
    slide = Presentation(output).slides[0]
    assert len(slide.shapes) == 4
    assert slide.shapes[0].text == 'Before bold both after & more'
    runs = slide.shapes[0].text_frame.paragraphs[0].runs
    assert all(run.font.size.pt == 24 for run in runs)
    assert all(str(run.font.color.rgb) == '123456' for run in runs)
    combined = next(run for run in runs if run.text == 'both')
    assert combined.font.bold and combined.font.italic
    assert slide.shapes[1].image.blob == (root / 'sample.png').read_bytes()
    table = slide.shapes[2].table
    assert table.cell(0, 0).span_height == 2
    assert table.cell(1, 0).is_spanned
    assert table.cell(0, 1).text == 'Top'
    assert table.cell(1, 1).text == 'Bottom'
    assert slide.shapes[3].image.size[0] > 0


@pytest.mark.parametrize('element', [
    {'elementType': 'icon', 'iconName': 'fas:lightbulb'},
    {'elementType': 'shape', 'shapeName': 'custom', 'viewBox': [10, 10], 'path': 'M0,0 L10,10'},
    {'elementType': 'shape', 'shapeName': 'rect', 'fill': {'type': 'gradient', 'color': '#fff'}},
    {'elementType': 'chart', 'series': [{'type': 'scatter', 'data': [1, 2]}]},
    {'elementType': 'image', 'src': 'sample.png', 'fit': {'mode': 'cover'}},
    {'elementType': 'text', 'content': {'text': '<sup>2</sup>'}},
    {'elementType': 'line', 'points': [[0, 0], [5, 5], [10, 0]]},
])
def test_unsupported_content_fails_without_replacing_output(tmp_path, element):
    root = make_deck(tmp_path, elements=[{'elementId': 'unsupported', 'bounds': [10, 10, 100, 100], **element}])
    Image.new('RGB', (100, 100), 'blue').save(root / 'sample.png')
    assert run_script('pptd_check.py', root).returncode != 0
    output = root / 'existing.pptx'
    Presentation().save(output)
    previous = output.read_bytes()
    result = run_script('pptd_to_pptx.py', root, '-o', str(output))
    assert result.returncode != 0
    assert 'unsupported' in result.stderr.lower()
    assert output.read_bytes() == previous


def test_missing_image_is_not_silently_omitted(tmp_path):
    root = make_deck(tmp_path, elements=[{'elementId': 'photo', 'elementType': 'image', 'bounds': [10, 10, 100, 100], 'src': 'missing.png'}])
    output = root / 'result.pptx'
    result = run_script('pptd_to_pptx.py', root, '-o', str(output))
    assert result.returncode != 0
    assert 'photo' in result.stderr and 'missing.png' in result.stderr
    assert not output.exists()


def test_unloaded_font_declaration_is_rejected(tmp_path):
    root = make_deck(tmp_path, customFonts=[{'family': 'Uninstalled', 'src': 'font.ttf'}])
    assert run_script('pptd_check.py', root).returncode != 0
    result = run_script('pptd_to_pptx.py', root)
    assert result.returncode != 0
    assert 'customFonts' in result.stderr
    assert not list(root.glob('*.pptx'))
