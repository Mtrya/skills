#!/usr/bin/env python3
"""pptd → pptx 导出器。

覆盖：text（含 strong/em 基础富文本、字号、颜色、对齐、行距）、shape（常用预置形状+纯色填充+边框）、
line、image（预处理后的 fill 图片）、table（合并、单元格填充、对齐）、chart（经 matplotlib 渲染为图片嵌入）。
不支持的功能报错停止，支持范围由 pptd_support 定义。

用法: python pptd_to_pptx.py <pptd目录> [-o out.pptx]
依赖: pip install python-pptx matplotlib pyyaml
"""
import argparse, html, os, re, sys, tempfile

from pptd_support import SHAPE_TYPES, prepare_deck, prepare_element, text_blocks

try:
    import yaml
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
except ImportError as e:
    sys.exit(f'缺依赖: {e}（pip install python-pptx pyyaml matplotlib）')

EMU_PER_PX = 9525   # 96dpi: 1px = 9525 EMU

SHAPES = {name: getattr(MSO_SHAPE, value) for name, value in SHAPE_TYPES.items()}

def rgb(c, colors):
    if c is None or c == '': return None
    if c.startswith('$'): c = colors[c[1:]]
    c = c.lstrip('#')
    if len(c) == 3:
        c = ''.join(part * 2 for part in c)
    return RGBColor.from_string(c)

def px(v):
    """按 96 dpi 将源画布坐标和尺寸换算为 EMU。"""
    return Emu(round(v * EMU_PER_PX))

def strip_tags(s):
    return re.sub(r'<[^>]+>', '', s).strip()

def rich_runs(para_el, text, base):
    """把 <strong>/<em> 标记的富文本拆成 run 写入段落。"""
    tokens = re.split(r'(</?(?:strong|em|u)>)', text)
    active = []
    for tok in tokens:
        if not tok: continue
        if tok.startswith('</'):
            active.pop(); continue
        if tok in ('<strong>', '<em>', '<u>'):
            active.append(tok[1:-1]); continue
        bold = 'strong' in active; italic = 'em' in active
        underline = 'u' in active
        r = para_el.add_run(); r.text = html.unescape(tok)
        f = r.font
        if base.get('fontSize'): f.size = Pt(base['fontSize'])
        if base.get('fontFamily'): f.name = base['fontFamily']
        color = base.get('color')
        if color is not None: f.color.rgb = color
        f.bold = bool(base.get('bold') or bold) or None
        f.italic = bool(base.get('italic') or italic) or None
        if underline: f.underline = True

def add_text(slide, el, W, H, colors):
    c = el.get('content') or {}
    x, y, w, h = el['bounds']
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame; tf.word_wrap = True
    base = {'fontSize': c.get('fontSize'), 'fontFamily': c.get('fontFamily'),
            'color': rgb(c.get('color',''), colors), 'bold': c.get('bold'), 'italic': c.get('italic')}
    # 文本样式已由 prepare_element 解析。
    raw = c.get('text','')
    blocks = text_blocks(raw)
    valign = (c.get('align') or {}).get('vertical','top')
    tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}.get(valign, MSO_ANCHOR.TOP)
    for i, blk in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        halign = (c.get('align') or {}).get('horizontal','left')
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT,
                       'justify': PP_ALIGN.JUSTIFY}.get(halign, PP_ALIGN.LEFT)
        if c.get('lineHeight'): p.line_spacing = c['lineHeight']
        rich_runs(p, blk, base)

def apply_fill_border(sp, el, colors):
    fill = el.get('fill') or {}
    if fill.get('type') == 'solid':
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill.get('color','#000000'), colors) or RGBColor(0,0,0)
    else:
        sp.fill.background()
    border = el.get('border') or {}
    if border.get('color'):
        sp.line.color.rgb = rgb(border['color'], colors) or RGBColor(0,0,0)
        sp.line.width = Pt(border.get('width', 1))
    else:
        sp.line.fill.background()

def add_shape(slide, el, W, H, colors):
    name = el.get('shapeName','rect')
    mso = SHAPES[name]
    x, y, w, h = el['bounds']
    sp = slide.shapes.add_shape(mso, px(x), px(y), px(w), px(h))
    apply_fill_border(sp, el, colors)
    sp.shadow.inherit = False

def add_line(slide, el, W, H, colors):
    x, y, w, h = el['bounds']
    pts = el.get('points') or [[0,0],[w,h]]
    x1, y1 = pts[0]; x2, y2 = pts[-1]
    conn = slide.shapes.add_connector(1, px(x+x1), px(y+y1), px(x+x2), px(y+y2))
    b = el.get('border') or {}
    if b.get('color'):
        conn.line.color.rgb = rgb(b['color'], colors) or RGBColor(0,0,0)
        conn.line.width = Pt(b.get('width',1))

def add_image(slide, el, W, H, root, tmpdir, idx):
    src = el.get('src','')
    path = src if src.startswith(('http://','https://')) else os.path.join(root, src)
    if src.startswith(('http://','https://')):
        import urllib.request
        path = os.path.join(tmpdir, f'image_{idx}')
        urllib.request.urlretrieve(src, path)
    if not os.path.isfile(path):
        raise FileNotFoundError(f'图片不存在: {src}')
    x, y, w, h = el['bounds']
    slide.shapes.add_picture(path, px(x), px(y), px(w), px(h))

def add_table(slide, el, W, H, colors):
    x, y, w, h = el['bounds']
    rows = el.get('rows') or []
    cw, rh = el.get('columnWidths'), el.get('rowHeights')
    ncol, nrow = len(cw or []), len(rh or rows)
    if not ncol or not nrow or len(rows) != nrow:
        raise ValueError('Table dimensions and rows must be nonempty and consistent')
    gt = slide.shapes.add_table(nrow, ncol, px(x), px(y), px(w), px(h)).table
    for j, frac in enumerate(cw): gt.columns[j].width = px(frac * w)
    for i, frac in enumerate(rh): gt.rows[i].height = px(frac * h)
    # 被合并区域覆盖的单元格在输入中省略。
    r = 0
    for row in rows:
        c = 0
        for celldef in row:
            while c < ncol and gt.cell(r, c).is_spanned: c += 1
            if c >= ncol:
                raise ValueError(f'Too many cells in table row {r + 1}')
            cell = gt.cell(r, c)
            cell.text = strip_tags(celldef.get('text',''))
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    if celldef.get('fontSize'): run.font.size = Pt(celldef['fontSize'])
                    if celldef.get('bold'): run.font.bold = True
                    col = rgb(celldef.get('color',''), colors)
                    if col is not None: run.font.color.rgb = col
                hal = (celldef.get('align') or {}).get('horizontal')
                if hal: p.alignment = {'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[hal]
            fillc = rgb((celldef.get('fill') or {}).get('color',''), colors)
            if fillc is not None:
                cell.fill.solid(); cell.fill.fore_color.rgb = fillc
            rs, cs = celldef.get('rowSpan',1), celldef.get('colSpan',1)
            if rs < 1 or cs < 1 or r + rs > nrow or c + cs > ncol:
                raise ValueError('Table merge extends beyond the table')
            if rs > 1 or cs > 1:
                cell.merge(gt.cell(r + rs - 1, c + cs - 1))
            c += cs
        r += 1

def add_chart(slide, el, W, H, colors, tmpdir, idx):
    """chart → matplotlib PNG 嵌入（近似渲染，保数据不保像素级样式）。"""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for fname in ('Noto Sans CJK SC', 'WenQuanYi Zen Hei', 'Microsoft YaHei', 'PingFang SC', 'SimHei'):
        if any(fname.lower() in f.name.lower() for f in font_manager.fontManager.ttflist):
            matplotlib.rcParams['font.sans-serif'] = [fname]
            matplotlib.rcParams['axes.unicode_minus'] = False
            break
    series = el.get('series') or []
    xdata = (el.get('xAxis') or {}).get('data') or []
    fig, ax = plt.subplots(figsize=(el['bounds'][2]/96, el['bounds'][3]/96), dpi=150)
    for s in series:
        t, d, name = s.get('type'), s.get('data') or [], s.get('name')
        col = s.get('color')
        if isinstance(col, str) and col.startswith('$'):
            col = colors.get(col[1:])
        if t == 'bar': ax.bar(xdata or range(len(d)), d, label=name, color=col)
        elif t in ('line','area'):
            ax.plot(xdata or range(len(d)), d, label=name, color=col, marker='o')
            if t == 'area': ax.fill_between(range(len(d)), d, alpha=.3, color=col)
        elif t == 'pie':
            ax.clear()
            ax.pie([v['value'] if isinstance(v, dict) else v for v in d],
                   labels=[v.get('name','') if isinstance(v, dict) else '' for v in d],
                   autopct='%1.1f%%')
        else:
            raise ValueError(f'Unsupported chart type: {t}')
    if (el.get('xAxis') or {}).get('name'): ax.set_xlabel(el['xAxis']['name'])
    if (el.get('yAxis') or {}).get('name'): ax.set_ylabel(el['yAxis']['name'])
    if el.get('title'): ax.set_title(strip_tags(str(el['title'].get('text',''))))
    if (el.get('legend') or {}).get('show') and any(s.get('name') for s in series): ax.legend()
    ax.spines[['top','right']].set_visible(False)
    path = os.path.join(tmpdir, f'chart_{idx}.png')
    fig.tight_layout(); fig.savefig(path, transparent=False); plt.close(fig)
    x, y, w, h = el['bounds']
    slide.shapes.add_picture(path, px(x), px(y), px(w), px(h))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('deck_dir'); ap.add_argument('-o', '--output', default=None)
    a = ap.parse_args()
    root = os.path.abspath(a.deck_dir)
    entry = [f for f in os.listdir(root) if f.endswith('.pptd')]
    if len(entry) != 1: sys.exit('作品目录必须包含且仅包含一个 .pptd 主入口')
    deck = yaml.safe_load(open(os.path.join(root, entry[0]), encoding='utf-8'))
    prepare_deck(deck)
    W, H = deck.get('size', [960, 540])
    theme = deck.get('theme') or {}
    colors = theme.get('colors') or {}

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory(prefix='pptd-export-') as tmpdir:
        render_pages(prs, blank, deck, root, W, H, theme, colors, tmpdir)

    out = a.output or os.path.join(root, os.path.basename(root) + '.pptx')
    prs.save(out)
    print(f'已导出: {out}')


def render_pages(prs, blank, deck, root, W, H, theme, colors, tmpdir):
    for i, rel in enumerate(deck.get('pages') or [], 1):
        page = yaml.safe_load(open(os.path.join(root, rel), encoding='utf-8'))
        slide = prs.slides.add_slide(blank)
        for j, el in enumerate(page.get('elements') or []):
            et = el.get('elementType')
            try:
                el = prepare_element(el, theme)
                if et == 'text': add_text(slide, el, W, H, colors)
                elif et == 'shape': add_shape(slide, el, W, H, colors)
                elif et == 'line': add_line(slide, el, W, H, colors)
                elif et == 'image': add_image(slide, el, W, H, root, tmpdir, f'{i}_{j}')
                elif et == 'table': add_table(slide, el, W, H, colors)
                elif et == 'chart': add_chart(slide, el, W, H, colors, tmpdir, f'{i}_{j}')
            except Exception as e:
                raise ValueError(f'page {i} element {el.get("elementId")}: {e}') from e

if __name__ == '__main__':
    main()
